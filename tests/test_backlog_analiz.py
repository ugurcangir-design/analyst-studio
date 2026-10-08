"""UAT Mutabakat — analiz durumu + takip hunisi + toplam denklemleri (offline, 0 token, Jira MOCK).

Kapsam:
  1. analiz_durumu: teknik / işlevsel / hata modu → yapildi · Amaç+Kriter / serbest metin → kismi · boş → yok
  2. takip_asamasi: kapandı · açıkta · teyit · analiz bekliyor/kısmi · atamaya hazır · geliştirmede
  3. mutabakat (Jira MOCK): denklemler kapanır (UAT, hedef, iptal, huni), bağ ≠ task, satır alanları
  4. Epic modunda kapsam dışından linkle gelen task toplama girmez, ayrı sayılır
  5. Excel: Özet sayfası + Takip/Analiz kolonları
"""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
os.environ.setdefault("BILDIRIM", "false")

from skills import backlog_senkron as B  # noqa: E402

_hata = 0


def chk(ad, kosul, ek=""):
    global _hata
    if not kosul:
        _hata += 1
    print(("  ✓ " if kosul else "  ✗ ") + ad + (f"  [{ek}]" if ek and not kosul else ""))


DOLU = "Bu bölümde yeterince uzun ve anlamlı bir içerik bulunuyor, en az altmış karakter olsun diye."
TEKNIK = f"\n1. Amaç ve Hedefler\n{DOLU}\n\n3. Teknik Gereksinimler\n{DOLU}\n\n11. Kabul Kriterleri\n• {DOLU}\n"
ISLEVSEL = (f"\nAmaç\n{DOLU}\n\nMevcut Durum\n{DOLU}\n\nBeklenen Davranış / İstenen Geliştirme\n{DOLU}\n\n"
            f"Kabul Kriterleri\n{DOLU}\n")
HATA = f"\nSorun\n{DOLU}\n\nÇözüm\n• {DOLU}\n"
KOPRU = f"\n📌 Orijinal Talep\nkısa talep\n\n🤖 Teknik Analiz\n\nAmaç\n{DOLU}\n\nÇözüm\n{DOLU}\n\nKabul Kriterleri\n{DOLU}\n"
KISMI = f"\nAmaç\n{DOLU}\n\nKabul Kriteri\n{DOLU}\n"

print("1) analiz_durumu")
for ad, metin, beklenen in [
    ("teknik şablon", TEKNIK, "yapildi"), ("işlevsel (mevcut/beklenen)", ISLEVSEL, "yapildi"),
    ("hata modu Sorun+Çözüm", HATA, "yapildi"), ("köprü (🤖 Teknik Analiz)", KOPRU, "yapildi"),
    ("Amaç + Kabul Kriteri", KISMI, "kismi"), ("şablonsuz uzun serbest metin", "x " * 500, "kismi"),
    ("boş", "", "yok"), ("kısa", "Tutar alanı düzeltilsin.", "yok"),
    ("başlık var ama içi boş", "\nAmaç\n\nTeknik Gereksinimler\nkısa\n\nKabul Kriterleri\n", "yok"),
]:
    a = B.analiz_durumu(metin)
    chk(f"{ad} → {beklenen}", a["seviye"] == beklenen, a)
chk("dolu başlık adları döner", B.analiz_durumu(TEKNIK)["dolu"] == ["amaç ve hedefler", "teknik gereksinimler", "kabul kriterleri"])
chk("uzun gövde satırı başlık sanılmaz", B.analiz_durumu("Amaç " + "x" * 100 + "\n" + DOLU)["dolu"] == [])

print("2) takip_asamasi")
H = lambda d, sv: {"durum": d, "analiz_seviye": sv}   # noqa: E731
for ad, args, beklenen in [
    ("UAT kapanmış", ("Tamam", [H("Backlog", "yok")], False), "kapandi"),
    ("task yok", ("Backlog", [], False), "acikta"),
    ("yalnız aday", ("Backlog", [], True), "teyit"),
    ("backlog'da analizsiz var", ("Backlog", [H("Backlog", "yapildi"), H("BACKLOG", "yok")], False), "analiz_bekliyor"),
    ("backlog kısmi", ("Backlog", [H("Backlog", "kismi"), H("Test", "yok")], False), "analiz_kismi"),
    ("backlog analizli", ("Backlog", [H("Backlog", "yapildi"), H("Devam Ediyor", "yok")], False), "atamaya_hazir"),
    ("backlog'da task kalmadı", ("Devam Ediyor", [H("Test", "yok")], False), "gelistirmede"),
]:
    chk(f"{ad} → {beklenen}", B.takip_asamasi(*args) == beklenen, B.takip_asamasi(*args))


print("3) mutabakat (Jira MOCK)")


def g(key, status="Backlog", desc="", links=(), tip="Görev", summary=None):
    return {"key": key, "summary": summary or f"{key} özet", "status": status, "type": tip, "assignee": "",
            "description": desc, "baglantililar": [{"key": k, "iliski": "relates to"} for k in links],
            "parent_key": "", "parent_type": "", "comments": []}


UAT = [
    g("UAT-1", links=["TR-1", "TR-2"]),         # 2 task: biri analizsiz backlog → analiz bekliyor
    g("UAT-2", links=["TR-3"]),                  # analizli backlog → atamaya hazır
    g("UAT-3", links=["TR-1"]),                  # TR-1'i paylaşır (bağ ≠ task)
    g("UAT-4"),                                  # açıkta
    g("UAT-5", status="Tamam"),                  # kapanmış (task yok)
    g("UAT-6", status="İptal Edildi"),           # iptal
    g("UAT-7", tip="Epic"),                      # kapsayıcı → hariç
    g("UAT-8", links=["TR-4"]),                  # task geliştirmede
]
HEDEF = [
    g("TR-1", desc=""), g("TR-2", desc=TEKNIK), g("TR-3", desc=TEKNIK), g("TR-4", status="Devam Ediyor"),
    g("TR-5", status="Tamam"), g("TR-6", status="Backlog"), g("TR-7", status="İptal Edildi"),
]
B._cloud_id = lambda: "cid"
B.jira_site_url = lambda c: "https://ornek.atlassian.net"
B._jql_ara = lambda jql, c: [dict(x) for x in UAT] if "UAT" in jql else []
B._hedef_gorevleri_topla = lambda mod, hp, hk, kw, c: [dict(x) for x in HEDEF]
B._keyleri_cek = lambda keys, c: []
d = B.mutabakat("UAT", ["TR"], "tum")
s = d["sayimlar"]
chk("UAT denklemi: toplam = eşleşen + aday + açıkta", s["uat_toplam"] == s["uat_eslesen"] + s["uat_aday"] + s["eslesmeyen_uat"] == 6, s)
chk("hedef denklemi: toplam = eşleşen + aday + eşleşmeyen", s["hedef_toplam"] == s["hedef_eslesen"] + s["hedef_aday"] + s["eslesmeyen_hedef"] == 6, s)
chk("iptal = UAT iptal + hedef iptal", s["iptal"] == s["uat_iptal"] + s["hedef_iptal"] == 2)
chk("Epic/Story hariç sayıldı", s["uat_kapsayici_haric"] == 1)
chk("bağ (5) ≠ eşleşen task (4) · eşleşen UAT (4)", (s["bag"], s["hedef_eslesen"], s["uat_eslesen"]) == (5, 4, 4), s)
chk("huni toplamı = UAT toplamı", sum(s["takip"].values()) == s["uat_toplam"], s["takip"])
kod = {r["uat_key"]: r["takip_kod"] for r in d["eslesenler"]}
kod.update({r["key"]: r["takip_kod"] for r in d["eslesmeyen_uat"]})
chk("aşamalar doğru", kod == {"UAT-1": "analiz_bekliyor", "UAT-2": "atamaya_hazir", "UAT-3": "analiz_bekliyor",
                              "UAT-4": "acikta", "UAT-5": "kapandi", "UAT-8": "gelistirmede"}, kod)
chk("backlog analiz (tekil task): TR-1 yok, TR-2/TR-3 yapıldı", s["backlog_analiz"] == {"yapildi": 2, "kismi": 0, "yok": 1}, s["backlog_analiz"])
chk("eşleşmeyen hedef aktif/tamam", (s["eslesmeyen_hedef_aktif"], s["eslesmeyen_hedef_tamam"]) == (1, 1), s)
r1 = [r for r in d["eslesenler"] if r["uat_key"] == "UAT-1"]
chk("satır alanları: analiz + takip + x/y analizli", r1 and all(r["uat_analizli"] == "1/2" for r in r1)
    and {r["hedef_analiz"] for r in r1} == {"Analiz yok", "Analiz yapıldı"} and r1[0]["takip"] == "Analiz bekliyor")
chk("eşleşmeyen hedefte analiz alanı", all("analiz" in r for r in d["eslesmeyen_hedef"]))

print("4) epic modu — kapsam dışı linkli task toplama girmez")
B._keyleri_cek = lambda keys, c: [g(k, desc=TEKNIK) for k in keys]
UAT.append(g("UAT-9", links=["TR-99"]))   # TR-99 epic kapsamı dışında
B._jql_ara = lambda jql, c: [dict(x) for x in UAT] if "UAT" in jql else []
d = B.mutabakat("UAT", ["TR"], "epic", hedef_keys=["TR-500"])
s = d["sayimlar"]
r99 = [r for r in d["eslesenler"] if r["hedef_key"] == "TR-99"]
chk("UAT-9 eşleşti (iş takip ediliyor) ama task 'kapsam dışı'", r99 and r99[0]["hedef_kapsam_disi"] is True)
chk("kapsam dışı ayrı sayıldı, hedef denklemi yine kapanır", s["kapsam_disi_hedef"] == 1
    and s["hedef_toplam"] == s["hedef_eslesen"] + s["hedef_aday"] + s["eslesmeyen_hedef"], s)

print("5) Excel")
from openpyxl import load_workbook  # noqa: E402
yol = B.rapor_uret(d, tempfile.mkdtemp())
wb = load_workbook(yol)
chk("sayfalar (Özet ilk)", wb.sheetnames[0] == "Özet" and "Eşleşenler" in wb.sheetnames, wb.sheetnames)
ozet = {str(r[0].value).strip(): r[1].value for r in wb["Özet"].iter_rows(min_row=2) if r[0].value}
chk("Özet: UAT toplam + huni + backlog analiz", ozet.get("UAT TOPLAM (iptal hariç)") == s["uat_toplam"]
    and ozet.get("Analiz bekliyor") == s["takip"]["analiz_bekliyor"] and ozet.get("Analiz yok") == s["backlog_analiz"]["yok"], ozet)
bas = [c.value for c in wb["Eşleşenler"][1]]
chk("Eşleşenler: Takip + Analiz + Dolu Başlıklar kolonları", {"Takip", "Analiz", "Dolu Başlıklar"} <= set(bas), bas)

print()
if _hata:
    print(f"BACKLOG ANALİZ TESTLERİ BAŞARISIZ ({_hata} hata)")
    sys.exit(1)
print("BACKLOG ANALİZ TESTLERİ GEÇTİ")
